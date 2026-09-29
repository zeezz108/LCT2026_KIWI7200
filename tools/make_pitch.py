"""Pitch deck (PPTX) from the experiment results and docs images.

usage: python tools/make_pitch.py [--out docs/pitch.pptx] [--team "Название команды"]
Numbers are read from data/results (offline runs of all bags) and data/results/synthetic/summary.json.
"""
import argparse
import collections
import json
from pathlib import Path

import numpy as np
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
IMG = ROOT / "docs" / "img"
RES = ROOT / "data" / "results"
BAGS = ["doubleT_obstacle", "doubleT_platform", "roundT_doubleT", "roundT_pressureGate_roundT",
        "roundT_squareT_pressureGate_squareT", "squareT_platform_squareT_switch"]

NAVY = RGBColor(0x1B, 0x2A, 0x41)
RED = RGBColor(0xD7, 0x26, 0x3D)
ORANGE = RGBColor(0xF4, 0x9D, 0x37)
GREEN = RGBColor(0x2E, 0x8B, 0x57)
GREY = RGBColor(0x55, 0x5B, 0x66)
LIGHT = RGBColor(0xF2, 0xF4, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


# ---------------------------------------------------------------- numbers
def load_rows(path):
    return [json.loads(l) for l in path.read_text().splitlines() if l.strip()]


def metrics():
    m = {"frames": 0, "empty_frames": 0, "empty_danger": 0, "empty_warning": 0}
    totals, p95s = [], []
    for b in BAGS:
        path = RES / f"{b}.jsonl"
        if not path.exists():
            continue
        rows = load_rows(path)
        lv = collections.Counter(r["level"] for r in rows)
        m["frames"] += len(rows)
        if b == "doubleT_obstacle":
            m["obstacle_danger_frames"] = lv[2]
            d = [o["distance"] for r in rows for o in r["obstacles"] if o["level"] == 2]
            m["obstacle_danger_distance"] = float(np.median(d)) if d else float("nan")
        else:
            m["empty_frames"] += len(rows)
            m["empty_danger"] += lv[2]
            m["empty_warning"] += lv[1]
        t = [r["ms"]["total"] for r in rows]
        totals += t
    m["danger_rate"] = 100.0 * m["empty_danger"] / max(1, m["empty_frames"])
    m["warning_rate"] = 100.0 * m["empty_warning"] / max(1, m["empty_frames"])
    m["ms_mean"] = float(np.mean(totals)) if totals else float("nan")
    m["ms_p95"] = float(np.percentile(totals, 95)) if totals else float("nan")
    syn = RES / "synthetic" / "summary.json"
    m["synthetic"] = json.loads(syn.read_text()) if syn.exists() else []
    return m


def recall_row(summary, obj, base="round"):
    rows = sorted((s for s in summary if s["object"] == obj and s["speed"] == 0 and s["lateral"] == 0
                   and s["name"].startswith(base)), key=lambda s: s["distance"])
    return [(s["distance"], 100 * s["recall_danger"]) for s in rows]


# ---------------------------------------------------------------- drawing helpers
def add_title(slide, prs, text, subtitle=None):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1.0))
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    tf = bar.text_frame
    tf.margin_left = Inches(0.5)
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(30)
    p.font.bold = True
    p.font.color.rgb = WHITE
    if subtitle:
        box = slide.shapes.add_textbox(Inches(0.5), Inches(1.05), prs.slide_width - Inches(1.0), Inches(0.5))
        q = box.text_frame.paragraphs[0]
        q.text = subtitle
        q.font.size = Pt(16)
        q.font.color.rgb = GREY


def add_bullets(slide, items, left, top, width, height, size=18):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        level = 0
        if isinstance(item, tuple):
            level, item = item
        p.text = ("• " if level == 0 else "– ") + item
        p.level = level
        p.font.size = Pt(size - 2 * level)
        p.font.color.rgb = NAVY if level == 0 else GREY
        p.space_after = Pt(6)
    return box


def add_number(slide, left, top, value, label, color, width=Inches(2.9)):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, Inches(1.7))
    card.fill.solid()
    card.fill.fore_color.rgb = LIGHT
    card.line.fill.background()
    tf = card.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = value
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = color
    p.alignment = PP_ALIGN.CENTER
    q = tf.add_paragraph()
    q.text = label
    q.font.size = Pt(13)
    q.font.color.rgb = GREY
    q.alignment = PP_ALIGN.CENTER


def add_image(slide, path, left, top, width=None, height=None):
    if not Path(path).exists():
        box = slide.shapes.add_textbox(left, top, width or Inches(4), Inches(0.5))
        box.text_frame.paragraphs[0].text = f"[нет изображения: {Path(path).name}]"
        return None
    return slide.shapes.add_picture(str(path), left, top, width=width, height=height)


# ---------------------------------------------------------------- slides
def build(out, team):
    m = metrics()
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    W = prs.slide_width

    # 1. title
    s = prs.slides.add_slide(blank)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY
    bg.line.fill.background()
    box = s.shapes.add_textbox(Inches(0.8), Inches(2.0), W - Inches(1.6), Inches(3))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Путь свободен?"
    p.font.size = Pt(54)
    p.font.bold = True
    p.font.color.rgb = WHITE
    for text, size, color in (("Обнаружение препятствий перед беспилотным поездом метро по 3D‑лидару", 26, WHITE),
                              ("ЛЦТ 2026 · кейс «Московский транспорт»", 20, ORANGE),
                              (team, 18, LIGHT)):
        q = tf.add_paragraph()
        q.text = text
        q.font.size = Pt(size)
        q.font.color.rgb = color
        q.space_before = Pt(12)

    # 2. task
    s = prs.slides.add_slide(blank)
    add_title(s, prs, "Задача", "Лидар Hesai 128 на голове поезда, ROS 2 Humble, Docker")
    add_bullets(s, [
        "Найти препятствие в габарите поезда впереди по ходу движения",
        "Сообщить: есть / нет, дистанция вдоль пути, время до столкновения",
        "Работать в реальном времени на незнакомых участках метро",
        "Критерии: качество и ложные тревоги, дальность 100 / 200 / 300 м, скорость, инженерия, запуск",
    ], Inches(0.6), Inches(1.7), Inches(7.2), Inches(4.5), size=22)
    add_image(s, IMG / "overview_doubleT_obstacle_f0.png", Inches(8.0), Inches(1.7), width=Inches(5.0))

    # 3. data insights
    s = prs.slides.add_slide(blank)
    add_title(s, prs, "Что показали данные", "6 записей, 2488 кадров; обучать нейросеть не на чем")
    add_bullets(s, [
        "Препятствия — только в одной записи: два человека, поезд стоит",
        "Разные крепления лидара: высота 1.34 и 1.73 м, крен 2.6°; разные топики и поле зрения",
        "Рельсы видны до 30–40 м, дальше путь не виден вообще",
        "Максимальная дальность ~200 м: на 150 м в человека попадает ~10 лучей, на 200 м — 3",
        "Дальние стены распадаются на редкие «полоски»: шаг Δx ≈ Δaz·x²/D (10 м на 120 м)",
        "→ опираемся на геометрию тоннеля, а не на обучение",
    ], Inches(0.6), Inches(1.7), Inches(12), Inches(5), size=22)

    # 4. idea
    s = prs.slides.add_slide(blank)
    add_title(s, prs, "Идея: три свойства любого тоннеля")
    cards = [("1", "Поезд едет по рельсам", "Шаблон двух рельсов в каждом кольце лидара даёт ось пути вблизи — в 100% кадров всех записей"),
             ("2", "Габарит проходит сквозь тоннель", "Изгиб и уклон вдали — тот, при котором «труба» габарита не задевает стен, а стены остаются на своих смещениях"),
             ("3", "Всё внутри трубы — препятствие", "Если возвышается над полотном и держится несколько кадров. Класс объекта не важен")]
    for i, (num, head, text) in enumerate(cards):
        left = Inches(0.6 + i * 4.2)
        card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.6), Inches(3.9), Inches(4.8))
        card.fill.solid()
        card.fill.fore_color.rgb = LIGHT
        card.line.fill.background()
        tf = card.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = num
        p.font.size = Pt(48)
        p.font.bold = True
        p.font.color.rgb = RED
        q = tf.add_paragraph()
        q.text = head
        q.font.size = Pt(22)
        q.font.bold = True
        q.font.color.rgb = NAVY
        r = tf.add_paragraph()
        r.text = text
        r.font.size = Pt(16)
        r.font.color.rgb = GREY
        r.space_before = Pt(10)

    # 5. pipeline
    s = prs.slides.add_slide(blank)
    add_title(s, prs, "Конвейер обработки кадра", "C++17, ~20 мс на кадр; ядро без зависимостей от ROS")
    steps = [("Калибровка", "плоскость полотна: высота, тангаж, крен"), ("Рельсы", "ось пути 3–32 м"),
             ("Коридор", "лучевой поиск кривизны и уклона до 200 м"), ("Зоны", "ОПАСНОСТЬ / ВНИМАНИЕ, отсев конструкций"),
             ("Трекинг", "подтверждение, скорость, TTC")]
    for i, (head, text) in enumerate(steps):
        left = Inches(0.4 + i * 2.55)
        shape = s.shapes.add_shape(MSO_SHAPE.CHEVRON if i else MSO_SHAPE.PENTAGON, left, Inches(1.7), Inches(2.6), Inches(1.3))
        shape.fill.solid()
        shape.fill.fore_color.rgb = NAVY if i % 2 == 0 else RGBColor(0x2E, 0x4A, 0x6B)
        shape.line.fill.background()
        p = shape.text_frame.paragraphs[0]
        p.text = head
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = WHITE
        box = s.shapes.add_textbox(left + Inches(0.2), Inches(3.05), Inches(2.3), Inches(0.9))
        box.text_frame.word_wrap = True
        q = box.text_frame.paragraphs[0]
        q.text = text
        q.font.size = Pt(13)
        q.font.color.rgb = GREY
    add_image(s, IMG / "pitch_corridor.png", Inches(0.4), Inches(4.0), width=Inches(12.5))

    # 6. key findings
    s = prs.slides.add_slide(blank)
    add_title(s, prs, "Что сделало решение устойчивым", "каждая находка — из разбора ошибок на данных")
    add_bullets(s, [
        "Дальняя ось: награда за стены на смещениях, измеренных у рельсов",
        (1, "ложные ОПАСНОСТИ вдали в двухпутном тоннеле исчезли; кривая R = 400 м отслеживается до 100 м"),
        "Lookahead по рельсам при отсечении гипотез",
        (1, "поиск терял верную кривую в 6 из 21 стартов → 0 из 21, без роста времени"),
        "Физика редких дальних стен: связывание ∝ x² только вне габарита",
        "Высота над реально видимым полом, плотность точек, «висящие» кластеры",
        "Автоматика для незнакомой записи: калибровка по полотну, топик, ось «вперёд», облака без ring",
    ], Inches(0.6), Inches(1.7), Inches(12), Inches(5.3), size=21)

    # 7. results: false alarms and real people
    s = prs.slides.add_slide(blank)
    add_title(s, prs, "Результаты: ложные тревоги", f"{m['empty_frames']} кадров записей без препятствий + запись с людьми")
    add_number(s, Inches(0.6), Inches(1.8), f"{m['danger_rate']:.1f}%", "кадров с ложной ОПАСНОСТЬЮ", GREEN)
    add_number(s, Inches(3.7), Inches(1.8), f"{m['warning_rate']:.1f}%", "кадров с ВНИМАНИЕМ\n(оборудование у стен)", ORANGE)
    add_number(s, Inches(6.8), Inches(1.8), f"{m.get('obstacle_danger_distance', float('nan')):.0f} м",
               "человек на пути обнаружен\n(реальная запись)", RED)
    add_number(s, Inches(9.9), Inches(1.8), f"{m['ms_mean']:.0f} мс", f"на кадр в среднем\np95 {m['ms_p95']:.0f} мс", NAVY)
    add_image(s, IMG / "exp_obstacle_timeline.png", Inches(3.2), Inches(3.65), height=Inches(3.75))

    # 8. results: range
    s = prs.slides.add_slide(blank)
    add_title(s, prs, "Результаты: дальность", "синтетические препятствия, трассированные лучами лидара в реальные записи")
    add_image(s, IMG / "exp_recall.png", Inches(0.4), Inches(1.6), width=Inches(7.2))
    rows = recall_row(m["synthetic"], "person")
    items = [f"Человек: {d:.0f} м — {r:.0f}%" for d, r in rows]
    approach = [x for x in m["synthetic"] if x["speed"] > 0]
    for a in approach:
        if a.get("first_danger_distance"):
            items.append(f"Сближение ({a['object']}): ОПАСНОСТЬ с {a['first_danger_distance']:.0f} м")
    items.append("Предел — физика: на 175–200 м в человека попадает 3–6 лучей")
    add_bullets(s, items, Inches(7.8), Inches(1.7), Inches(5.2), Inches(5.3), size=17)

    # 9. engineering
    s = prs.slides.add_slide(blank)
    add_title(s, prs, "Инженерия и запуск")
    add_bullets(s, [
        "docker build → docker run → ros2 bag play → статус в топике и в логе раз в секунду",
        "docker compose --profile demo: детектор + RViz2 + проигрывание записи одной командой",
        "25 юнит‑тестов на синтетическом тоннеле с трассировкой лучей; прогоняются при сборке образа",
        "Офлайн‑раннер: запись обрабатывается быстрее реального времени, JSONL для анализа",
        "Инструменты оценки: генератор препятствий, серия сценариев, отладочные кадры",
        "Выход: ObstacleStatus, vision_msgs/Detection3DArray, маркеры и облако по зонам для RViz2",
        "Документация: архитектура, алгоритм, эксперименты, журнал исследований",
    ], Inches(0.6), Inches(1.5), Inches(7.4), Inches(5.5), size=19)
    add_image(s, IMG / "rviz_danger.png", Inches(8.2), Inches(1.5), width=Inches(4.8))

    # 10. limits and next steps
    s = prs.slides.add_slide(blank)
    add_title(s, prs, "Ограничения и развитие")
    add_bullets(s, [
        "Сейчас",
        (1, "~150 м для человека — предел плотности лучей; 300 м недостижимы этим сенсором"),
        (1, "на кривой видно только до стены; при смене сечения тоннеля ось вдали держится хуже"),
        (1, "низкие объекты (< 0.4 м) вдали неотличимы от элементов пути"),
        "Дальше",
        (1, "карта пути (априорная ось и габарит по пикетажу) → точная зона на всей дальности"),
        (1, "слияние с камерой и радаром для дальних и низких объектов"),
        (1, "накопление реальных случаев → классификатор поверх геометрии"),
    ], Inches(0.6), Inches(1.5), Inches(12), Inches(5.5), size=21)

    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    print("saved", out)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=str(ROOT / "docs" / "pitch.pptx"))
    ap.add_argument("--team", default="Команда: [название]")
    args = ap.parse_args()
    build(Path(args.out), args.team)


if __name__ == "__main__":
    main()
